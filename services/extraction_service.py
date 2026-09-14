"""Text extraction with explicit status and OCR fallback reporting."""

from __future__ import annotations
import logging
from pathlib import Path

from config import settings

logger = logging.getLogger(__name__)

if settings.TESSERACT_CMD:
    import pytesseract
    pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD


def extract_text(file_path: Path) -> str:
    return extract_text_detailed(file_path)["text"]


def extract_text_detailed(file_path: Path) -> dict[str, str | None]:
    extension = file_path.suffix.lower()
    try:
        if extension == ".pdf":
            text = _extract_pdf(file_path)
        elif extension == ".docx":
            text = _extract_docx(file_path)
        elif extension == ".xlsx":
            text = _extract_xlsx(file_path)
        elif extension in {".jpg", ".jpeg", ".png"}:
            text = _extract_image(file_path)
        else:
            return {"text": "", "status": "failed", "error": "Unsupported file type", "source": None}
        if text.strip():
            return {"text": text.strip(), "status": "success", "error": None, "source": extension[1:]}
        return {"text": "", "status": "failed", "error": "No text could be extracted", "source": extension[1:]}
    except Exception as exc:
        logger.exception("Text extraction failed for %s", file_path)
        return {"text": "", "status": "failed", "error": str(exc)[:500], "source": extension[1:]}


def _extract_pdf(path: Path) -> str:
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages).strip()
    if len(text) >= 50:
        return text
    try:
        import fitz
        with fitz.open(str(path)) as doc:
            fallback = "\n".join(page.get_text("text") for page in doc).strip()
        if len(fallback) > len(text):
            text = fallback
    except Exception as exc:
        logger.warning("PDF text fallback failed: %s", exc)
    if len(text) < 50:
        try:
            import fitz
            from PIL import Image
            with fitz.open(str(path)) as doc:
                ocr = []
                for page in doc:
                    pixmap = page.get_pixmap(dpi=200)
                    image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                    ocr.append(_ocr(image))
            text = "\n".join(ocr).strip() or text
        except Exception as exc:
            logger.warning("PDF OCR fallback failed: %s", exc)
    return text


def _extract_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    parts = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    parts.extend(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()) for table in doc.tables for row in table.rows)
    return "\n".join(part for part in parts if part).strip()


def _extract_xlsx(path: Path) -> str:
    import pandas as pd
    sheets = pd.read_excel(path, sheet_name=None, dtype=str, engine="openpyxl")
    return "\n".join(f"Sheet: {name}\n{frame.fillna('').to_csv(index=False)}" for name, frame in sheets.items()).strip()


def _extract_image(path: Path) -> str:
    from PIL import Image
    return _ocr(Image.open(path)).strip()


def _ocr(image) -> str:
    import pytesseract
    try:
        return pytesseract.image_to_string(image, lang=settings.OCR_LANGUAGE)
    except Exception as exc:
        logger.warning("OCR failed with language %s: %s", settings.OCR_LANGUAGE, exc)
        if settings.OCR_LANGUAGE != "eng":
            try:
                return pytesseract.image_to_string(image, lang="eng")
            except Exception as fallback_exc:
                logger.warning("English OCR fallback failed: %s", fallback_exc)
        return ""
